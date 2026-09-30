# rotation_planner.py
"""
Standalone N-month rotation planner.

WHAT IT DOES
------------
  1. Decides which station each doctor is assigned to for the next N months,
     with a HARD constraint that each doctor visits at least K distinct
     stations across those N months.
  2. Distributes vacation days across those months.
  3. Optionally writes one monthly template file per month, in the SAME
     structure as your existing Stationsplan templates.

WHAT IT DOES NOT DO
-------------------
  - It does NOT assign daily duties (SD, ZD, KM, ...). That is the job of
    the existing app.
  - It does NOT touch the existing Rules_updated.xlsx used by the app.
    All planner output goes to a separate file (default: RotationPlan.xlsx).

CONFIGURATION (optional sheets in Rules.xlsx)
---------------------------------------------
  RotationConfig:
      | Key                       | Value                        |
      | PlanStartMonth            | 2026-10                      |
      | PlanMonths                | 12                           |
      | MinDistinctStations       | 2                            |
      | VacationDaysPerDoctor     | 30                           |
      | GenerateTemplateFiles     | Yes / No                     |
      | TemplatePath              | /path/to/templates           |
      | TemplatePattern           | Stationsplan {month_name} {yy}.xlsx |

  AllowedStations (optional, strongly recommended):
      | Doctor | AllowedStations                          |
      | Jann   | 65 PP;65 LAF;85 Häm/Onk/Rheu             |
      | Rieber | 65 PP;65 LAF                             |
    If a doctor is not listed, all stations are allowed for them.
    If a doctor is listed with fewer than MinDistinctStations stations,
    the planner will report an error instead of silently relaxing.

  ResearchFree (optional):
      | Doctor | Month   | Capacity |
      | Jann   | 2026-12 | 50       |

  VacationRules (optional):
      | Doctor | PreferredMonths |
      | Jann   | 2026-11;2027-02 |

USAGE
-----
  CLI:
      python rotation_planner.py --rules Rules.xlsx --out RotationPlan.xlsx
      python rotation_planner.py --rules Rules.xlsx --start 2026-10 --months 12
      python rotation_planner.py --rules Rules.xlsx --generate-templates

  Python:
      from rotation_planner import plan_rotation
      plan_rotation("Rules.xlsx", "RotationPlan.xlsx", start_month="2026-10")
"""

from __future__ import annotations

import argparse
import os
import re
from calendar import monthrange
from datetime import datetime
from typing import Dict, List, Optional, Tuple

import pandas as pd
from ortools.sat.python import cp_model


# ----------------------------------------------------------------------
# Constants
# ----------------------------------------------------------------------
ROTATION_SHEET = "RotationPlan"
VACATION_SHEET = "VacationPlan"
RESEARCH_FREE_SHEET = "ResearchFreePlan"

DEFAULT_CONFIG = {
    "PlanStartMonth": None,
    "PlanMonths": 6,
    "MinDistinctStations": 2,        # NEW: hard constraint
    "VacationDaysPerDoctor": 30,
    "ResearchFreeDaysPerDoctor": 0,
    "GenerateTemplateFiles": "No",
    "TemplatePath": "",
    "TemplatePattern": "Stationsplan {month_name} {yy}.xlsx",
}

GERMAN_MONTHS = {
    1: "Januar", 2: "Februar", 3: "März", 4: "April",
    5: "Mai", 6: "Juni", 7: "Juli", 8: "August",
    9: "September", 10: "Oktober", 11: "November", 12: "Dezember",
}


# ----------------------------------------------------------------------
# Sheet loaders
# ----------------------------------------------------------------------
def _read_sheet(path: str, sheet: str) -> Optional[pd.DataFrame]:
    try:
        return pd.read_excel(path, sheet_name=sheet)
    except Exception:
        return None


def _load_config(rules_path: str) -> Dict[str, object]:
    cfg = dict(DEFAULT_CONFIG)
    df = _read_sheet(rules_path, "RotationConfig")
    if df is None or df.empty:
        return cfg
    df = df.dropna(subset=["Key"])
    for _, row in df.iterrows():
        key = str(row["Key"]).strip()
        val = row["Value"]
        if key in cfg and not pd.isna(val):
            cfg[key] = val
    return cfg


def _load_doctors(rules_path: str) -> pd.DataFrame:
    df = _read_sheet(rules_path, "Doctors")
    if df is None or df.empty:
        raise ValueError("Rules file has no Doctors sheet.")
    df = df.copy()
    df["Name"] = df["Name"].astype(str).str.strip()
    df = df[df["Name"] != ""]
    if "FTE (%)" not in df.columns and "FTE" in df.columns:
        df["FTE (%)"] = df["FTE"]
    df["FTE (%)"] = pd.to_numeric(df["FTE (%)"], errors="coerce").fillna(100)
    df["Station"] = df["Station"].astype(str).str.strip()
    if "Active" not in df.columns:
        df["Active"] = "Yes"
    return df


def _load_stations(rules_path: str) -> pd.DataFrame:
    df = _read_sheet(rules_path, "Stations")
    if df is None or df.empty:
        raise ValueError("Rules file has no Stations sheet.")
    df = df.copy()
    df["Station"] = df["Station"].astype(str).str.strip()
    return df


def _load_allowed_stations(rules_path: str) -> pd.DataFrame:
    """
    Optional sheet:
        | Doctor | AllowedStations                    |
        | Jann   | 65 PP;65 LAF;85 Häm/Onk/Rheu       |
    Returns empty DataFrame if sheet is missing.
    """
    df = _read_sheet(rules_path, "AllowedStations")
    if df is None:
        return pd.DataFrame(columns=["Doctor", "AllowedStations"])
    df = df.copy()
    df["Doctor"] = df["Doctor"].astype(str).str.strip()
    df["AllowedStations"] = df.get("AllowedStations", "").astype(str)
    return df


def _load_research_free(rules_path: str) -> pd.DataFrame:
    df = _read_sheet(rules_path, "ResearchFree")
    if df is None:
        return pd.DataFrame(columns=["Doctor", "Month", "Capacity"])
    df = df.copy()
    df["Doctor"] = df["Doctor"].astype(str).str.strip()
    df["Month"] = df["Month"].astype(str).str.strip()
    df["Capacity"] = pd.to_numeric(df["Capacity"], errors="coerce").fillna(0)
    return df


def _load_vacation_rules(rules_path: str) -> pd.DataFrame:
    df = _read_sheet(rules_path, "VacationRules")
    if df is None:
        return pd.DataFrame(columns=["Doctor", "PreferredMonths"])
    df = df.copy()
    df["Doctor"] = df["Doctor"].astype(str).str.strip()
    df["PreferredMonths"] = df.get("PreferredMonths", "").astype(str)
    return df


# ----------------------------------------------------------------------
# Month utilities
# ----------------------------------------------------------------------
def _month_range(start: datetime, n: int) -> List[str]:
    out = []
    y, m = start.year, start.month
    for _ in range(n):
        out.append(f"{y:04d}-{m:02d}")
        m += 1
        if m > 12:
            m = 1
            y += 1
    return out


def _month_days(month_str: str) -> int:
    y, m = map(int, month_str.split("-"))
    return monthrange(y, m)[1]


def _month_label(month_str: str) -> Tuple[str, str]:
    y, m = month_str.split("-")
    return GERMAN_MONTHS[int(m)], y[2:]


# ----------------------------------------------------------------------
# Rotation solver (CP-SAT) — with MinDistinctStations hard constraint
# ----------------------------------------------------------------------
def _build_allowed_station_map(
    doctors: List[str],
    stations: List[str],
    allowed_df: pd.DataFrame,
    min_distinct: int,
) -> Tuple[Dict[str, List[str]], List[str]]:
    """
    Returns:
        allowed_map: doctor -> list of allowed stations (subset of stations)
        errors:      list of error strings (empty if OK)
    """
    errors: List[str] = []
    allowed_map: Dict[str, List[str]] = {}

    allowed_lookup: Dict[str, List[str]] = {}
    for _, row in allowed_df.iterrows():
        doc = str(row["Doctor"]).strip()
        raw = str(row["AllowedStations"]).strip()
        if not doc or raw.lower() == "nan" or not raw:
            continue
        allowed_lookup[doc] = [s.strip() for s in raw.split(";") if s.strip()]

    for d in doctors:
        if d in allowed_lookup:
            # Filter to stations that actually exist in Stations sheet
            allowed = [s for s in allowed_lookup[d] if s in stations]
            if not allowed:
                errors.append(
                    f"Doctor '{d}' has AllowedStations entry but none of the "
                    f"listed stations exist in Stations sheet: {allowed_lookup[d]}"
                )
                allowed = list(stations)
            elif len(allowed) < min_distinct:
                errors.append(
                    f"Doctor '{d}' is allowed only {len(allowed)} station(s) "
                    f"({allowed}), but MinDistinctStations = {min_distinct}. "
                    f"Either lower MinDistinctStations, or add more stations "
                    f"to AllowedStations for this doctor."
                )
        else:
            allowed = list(stations)
        allowed_map[d] = allowed

    return allowed_map, errors


def _solve_rotation(
    doctors: List[str],
    stations: List[str],
    months: List[str],
    doctors_df: pd.DataFrame,
    stations_df: pd.DataFrame,
    allowed_map: Dict[str, List[str]],
    min_distinct: int,
    verbose: bool = True,
) -> pd.DataFrame:
    """
    Decide x[doctor, station, month] ∈ {0,1}.

    Hard constraints:
      - Each doctor: exactly one station per month.
      - Each station: exactly `RequiredDoctors` doctors per month.
      - Each doctor: visits at least `min_distinct` distinct stations
        across the N months.
      - Each doctor: only stations from allowed_map[doctor].

    Soft objective:
      - Konstanz: reward staying at the same station month-to-month.
      - Home station: reward being at the doctor's home station.
    """
    model = cp_model.CpModel()

    # ---- Vars ----
    x: Dict[Tuple[str, str, str], cp_model.IntVar] = {}
    for d in doctors:
        for s in allowed_map[d]:
            for m in months:
                x[(d, s, m)] = model.NewBoolVar(f"x_{d}_{s}_{m}")

    # ---- 1. Each doctor exactly one station per month (only allowed ones) ----
    for d in doctors:
        for m in months:
            model.Add(sum(x[(d, s, m)] for s in allowed_map[d]) == 1)

    # ---- 2. Station headcount per month ----
    station_required: Dict[str, int] = {}
    has_required = "RequiredDoctors" in stations_df.columns
    for _, row in stations_df.iterrows():
        s = str(row["Station"]).strip()
        if has_required:
            try:
                station_required[s] = int(row["RequiredDoctors"])
            except Exception:
                station_required[s] = 0
        else:
            station_required[s] = int((doctors_df["Station"] == s).sum())

    for s in stations:
        req = station_required.get(s, 0)
        for m in months:
            if req == 0:
                for d in doctors:
                    if s in allowed_map[d]:
                        model.Add(x[(d, s, m)] == 0)
            else:
                terms = [x[(d, s, m)] for d in doctors if s in allowed_map[d]]
                if terms:
                    model.Add(sum(terms) == req)
                elif req > 0:
                    # No doctor is allowed to be here, but station requires staff
                    # → this is infeasible, let CP-SAT report it
                    model.Add(0 == 1)

    # ---- 3. HARD: each doctor visits at least K distinct stations ----
    ever_in: Dict[Tuple[str, str], cp_model.IntVar] = {}
    for d in doctors:
        # For each allowed station, a boolean: did the doctor ever work there?
        station_bools = []
        for s in allowed_map[d]:
            b = model.NewBoolVar(f"ever_{d}_{s}")
            # b == 1  ⇔  sum over months of x[d,s,m] >= 1
            # Implement with a reified "at_least_one_month_here"
            n_months_here = model.NewIntVar(0, len(months), f"nmonths_{d}_{s}")
            model.Add(n_months_here == sum(x[(d, s, m)] for m in months))
            model.Add(b == 1).OnlyEnforceIf(n_months_here >= 1) if False else None
            # Use a standard trick: b >= (n_months_here / len(months))
            # Simpler: enforce b == 1 iff n_months_here >= 1, via two directions
            #  - If n_months_here >= 1, then b == 1
            #  - If n_months_here == 0, then b == 0
            # CP-SAT supports Add(n_months_here >= 1).OnlyEnforceIf(b)
            # and Add(n_months_here == 0).OnlyEnforceIf(b.Not())
            model.Add(n_months_here >= 1).OnlyEnforceIf(b)
            model.Add(n_months_here == 0).OnlyEnforceIf(b.Not())
            ever_in[(d, s)] = b
            station_bools.append(b)

        # At least min_distinct of these are 1
        model.Add(sum(station_bools) >= min_distinct)

    # ---- Soft objective ----
    objective = []

    # Konstanz reward (soft): staying at the same station month-to-month
    KONSTANZ_REWARD = 3   # lowered from 10 so that rotation is not blocked
    for d in doctors:
        for s in allowed_map[d]:
            for i in range(1, len(months)):
                m0, m1 = months[i - 1], months[i]
                stay = model.NewBoolVar(f"stay_{d}_{s}_{i}")
                model.Add(stay <= x[(d, s, m0)])
                model.Add(stay <= x[(d, s, m1)])
                model.Add(stay >= x[(d, s, m0)] + x[(d, s, m1)] - 1)
                objective.append(-KONSTANZ_REWARD * stay)

    # Home-station reward (soft)
    home = {row["Name"]: str(row.get("Station", "")).strip()
            for _, row in doctors_df.iterrows()}
    HOME_REWARD = 2
    for d in doctors:
        home_s = home.get(d, "")
        if home_s and home_s in allowed_map[d]:
            for m in months:
                objective.append(-HOME_REWARD * x[(d, home_s, m)])

    model.Minimize(sum(objective))

    # ---- Solve ----
    solver = cp_model.CpSolver()
    solver.parameters.max_time_in_seconds = 120.0
    solver.parameters.num_search_workers = 4
    status = solver.Solve(model)

    if status not in (cp_model.OPTIMAL, cp_model.FEASIBLE):
        # Provide a diagnostic summary
        diag_lines = [
            "Rotation planner is INFEASIBLE.",
            "",
            "Common causes:",
            f"  - MinDistinctStations = {min_distinct} is too high for some doctors.",
            "    Each doctor needs at least this many allowed stations.",
            "  - RequiredDoctors sums do not match the number of doctors.",
            "    Each station must be staffed exactly as RequiredDoctors says.",
            "  - AllowedStations (if provided) restrict doctors too much.",
            "",
            "Doctors and their allowed stations:",
        ]
        for d in doctors[:20]:
            diag_lines.append(f"  {d}: {allowed_map[d]}")
        if len(doctors) > 20:
            diag_lines.append(f"  ... and {len(doctors) - 20} more.")
        raise RuntimeError("\n".join(diag_lines))

    rows = []
    for d in doctors:
        row = {"Doctor": d}
        for m in months:
            for s in allowed_map[d]:
                if solver.Value(x[(d, s, m)]) == 1:
                    row[m] = s
                    break
            else:
                row[m] = ""
        rows.append(row)

    if verbose:
        n_distinct_per_doc = []
        for _, row in pd.DataFrame(rows).iterrows():
            vals = [row[m] for m in months]
            n_distinct_per_doc.append(len(set(v for v in vals if v)))
        print(
            f"[rotation] solved {len(doctors)} doctors × {len(months)} months; "
            f"min distinct stations per doctor: "
            f"{min(n_distinct_per_doc) if n_distinct_per_doc else 0}"
        )

    return pd.DataFrame(rows, columns=["Doctor"] + months)


# ----------------------------------------------------------------------
# Vacation assignment (CP-SAT)
# ----------------------------------------------------------------------
def _solve_vacation(
    doctors: List[str],
    months: List[str],
    total_days: int,
    preferred: Dict[str, List[str]],
    verbose: bool = True,
) -> pd.DataFrame:
    model = cp_model.CpModel()
    max_per_month = {m: int(_month_days(m) * 0.4) for m in months}

    v: Dict[Tuple[str, str], cp_model.IntVar] = {}
    for d in doctors:
        for m in months:
            v[(d, m)] = model.NewIntVar(0, max_per_month[m], f"v_{d}_{m}")

    for d in doctors:
        model.Add(sum(v[(d, m)] for m in months) == total_days)

    objective = []
    for d in doctors:
        for m in preferred.get(d, []):
            if m in months:
                objective.append(-v[(d, m)])

    avg = total_days / len(months)
    avg_int = int(round(avg))
    for d in doctors:
        for m in months:
            dev = model.NewIntVar(0, max_per_month[m], f"dev_{d}_{m}")
            pos = model.NewIntVar(0, max_per_month[m], f"pos_{d}_{m}")
            neg = model.NewIntVar(0, max_per_month[m], f"neg_{d}_{m}")
            model.Add(v[(d, m)] - avg_int == pos - neg)
            model.Add(dev == pos + neg)
            objective.append(2 * dev)

    model.Minimize(sum(objective))

    solver = cp_model.CpSolver()
    solver.parameters.max_time_in_seconds = 30.0
    solver.parameters.num_search_workers = 4
    status = solver.Solve(model)

    if status not in (cp_model.OPTIMAL, cp_model.FEASIBLE):
        raise RuntimeError(f"Vacation planner infeasible. Status: {status}")

    rows = []
    for d in doctors:
        row = {"Doctor": d}
        for m in months:
            row[m] = solver.Value(v[(d, m)])
        rows.append(row)

    if verbose:
        print(f"[vacation] distributed {total_days} days per doctor")

    return pd.DataFrame(rows, columns=["Doctor"] + months)


# ----------------------------------------------------------------------
# Template generation
# ----------------------------------------------------------------------
def _pattern_to_regex(pattern: str) -> str:
    regex = re.escape(pattern)
    regex = regex.replace(re.escape("{month_name}"), r"(?P<month_name>[A-Za-zÄÖÜäöüß]+)")
    regex = regex.replace(re.escape("{yy}"), r"(?P<yy>\d{2})")
    regex = regex.replace(re.escape("{yyyy}"), r"(?P<yyyy>\d{4})")
    regex = regex.replace(re.escape("{mm}"), r"(?P<mm>\d{2})")
    return regex


def _find_template_for_month(
    template_dir: str,
    pattern: str,
    month_str: str,
) -> Optional[str]:
    if not template_dir or not os.path.isdir(template_dir):
        return None
    month_name, yy = _month_label(month_str)
    regex = _pattern_to_regex(pattern)
    compiled = re.compile(regex, re.IGNORECASE)
    for fname in os.listdir(template_dir):
        m = compiled.match(fname)
        if not m:
            continue
        try:
            if m.groupdict().get("month_name", "").lower() == month_name.lower() \
               and m.groupdict().get("yy", "") == yy:
                return os.path.join(template_dir, fname)
        except IndexError:
            continue
    return None


def _rearrange_template(
    template_path: str,
    rotation_df: pd.DataFrame,
    month_str: str,
    doctors_df: pd.DataFrame,
    stations_df: pd.DataFrame,
    output_path: str,
) -> Optional[str]:
    import openpyxl

    if rotation_df is None or rotation_df.empty:
        return None
    if month_str not in rotation_df.columns:
        print(f"[template] no rotation info for {month_str}, skipping")
        return None

    desired: Dict[str, List[str]] = {s: [] for s in stations_df["Station"].tolist()}
    for _, row in rotation_df.iterrows():
        doc = str(row["Doctor"]).strip()
        st = str(row.get(month_str, "")).strip()
        if doc and st and st in desired:
            desired[st].append(doc)

    wb = openpyxl.load_workbook(template_path)
    ws = wb.active

    doctor_name_col = 1
    doctor_rows: Dict[str, Tuple[int, Dict[int, object]]] = {}
    for r in range(2, ws.max_row + 1):
        name_cell = ws.cell(row=r, column=doctor_name_col).value
        if not name_cell or not isinstance(name_cell, str):
            continue
        name = name_cell.strip().rstrip("*").strip()
        if name in doctors_df["Name"].tolist():
            values = {c: ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)}
            doctor_rows[name] = (r, values)

    station_names = set(stations_df["Station"].tolist())
    station_positions: Dict[str, int] = {}
    for r in range(2, ws.max_row + 1):
        name_cell = ws.cell(row=r, column=doctor_name_col).value
        if not name_cell or not isinstance(name_cell, str):
            continue
        name = name_cell.strip()
        if name in station_names:
            station_positions[name] = r

    ordered_stations = sorted(station_positions.items(), key=lambda x: x[1])

    for idx, (station, row) in enumerate(ordered_stations):
        next_station_row = (
            ordered_stations[idx + 1][1] if idx + 1 < len(ordered_stations) else ws.max_row + 1
        )
        block_rows = list(range(row + 1, next_station_row))
        doctor_block_rows = [
            r for r in block_rows
            if ws.cell(row=r, column=doctor_name_col).value
            and isinstance(ws.cell(row=r, column=doctor_name_col).value, str)
            and ws.cell(row=r, column=doctor_name_col).value.strip().rstrip("*").strip()
            in doctors_df["Name"].tolist()
        ]

        assigned = desired.get(station, [])
        for i, r in enumerate(doctor_block_rows):
            new_name = assigned[i] if i < len(assigned) else ""
            if new_name:
                ws.cell(row=r, column=doctor_name_col).value = new_name
            else:
                ws.cell(row=r, column=doctor_name_col).value = None

    wb.save(output_path)
    return output_path


# ----------------------------------------------------------------------
# Main entry point
# ----------------------------------------------------------------------
def plan_rotation(
    rules_path: str = "Rules.xlsx",
    output_path: str = "RotationPlan.xlsx",
    start_month: Optional[str] = None,
    n_months: Optional[int] = None,
    min_distinct: Optional[int] = None,
    generate_templates: Optional[bool] = None,
    verbose: bool = True,
) -> Dict[str, pd.DataFrame]:
    """
    Run the full planner.

    Returns a dict with keys: 'rotation', 'vacation', 'research_free'.
    """
    cfg = _load_config(rules_path)

    if start_month is None:
        start_month = cfg.get("PlanStartMonth")
    if n_months is None:
        n_months = int(cfg.get("PlanMonths", 6))
    if min_distinct is None:
        min_distinct = int(cfg.get("MinDistinctStations", 2))
    if generate_templates is None:
        generate_templates = str(cfg.get("GenerateTemplateFiles", "No")).lower() == "yes"

    if start_month is None:
        today = datetime.today()
        y, m = today.year, today.month + 1
        if m > 12:
            m = 1
            y += 1
        start_month = f"{y:04d}-{m:02d}"

    months = _month_range(datetime.strptime(start_month + "-01", "%Y-%m-%d"), n_months)

    if verbose:
        print(f"[planner] planning {n_months} months: {months[0]} → {months[-1]}")
        print(f"[planner] MinDistinctStations = {min_distinct}")

    doctors_df = _load_doctors(rules_path)
    stations_df = _load_stations(rules_path)
    allowed_df = _load_allowed_stations(rules_path)
    research_free_df = _load_research_free(rules_path)
    vacation_rules_df = _load_vacation_rules(rules_path)

    doctors = [str(d).strip() for d in doctors_df["Name"].tolist()
               if str(d).strip() and str(d).strip() != "nan"]
    stations = [str(s).strip() for s in stations_df["Station"].tolist()
                if str(s).strip() and str(s).strip() != "nan"]

    # Build allowed-station map, report obvious conflicts early
    allowed_map, errors = _build_allowed_station_map(
        doctors, stations, allowed_df, min_distinct
    )
    if errors:
        msg = "\n".join(errors)
        raise ValueError(
            f"Configuration errors found:\n{msg}\n\n"
            f"Fix the AllowedStations sheet or lower MinDistinctStations."
        )

    # ----- 1. Rotation -----
    rotation_df = _solve_rotation(
        doctors, stations, months, doctors_df, stations_df,
        allowed_map, min_distinct, verbose,
    )

    # ----- 2. Vacation -----
    preferred: Dict[str, List[str]] = {}
    for _, row in vacation_rules_df.iterrows():
        doc = str(row["Doctor"]).strip()
        months_str = str(row.get("PreferredMonths", "")).strip()
        if doc and months_str and months_str.lower() != "nan":
            preferred[doc] = [m.strip() for m in months_str.split(";") if m.strip()]

    total_vac_days = int(cfg.get("VacationDaysPerDoctor", 30))
    vacation_df = _solve_vacation(doctors, months, total_vac_days, preferred, verbose)

    # ----- 3. Research-free -----
    research_free_plan_rows = []
    for _, row in research_free_df.iterrows():
        doc = str(row["Doctor"]).strip()
        month = str(row["Month"]).strip()
        capacity = float(row["Capacity"])
        if doc in doctors and month in months:
            research_free_plan_rows.append({
                "Doctor": doc,
                "Month": month,
                "Capacity": capacity,
            })
    research_free_plan_df = pd.DataFrame(
        research_free_plan_rows,
        columns=["Doctor", "Month", "Capacity"],
    )

    # ----- 4. Write output -----
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        rotation_df.to_excel(writer, sheet_name=ROTATION_SHEET, index=False)
        vacation_df.to_excel(writer, sheet_name=VACATION_SHEET, index=False)
        research_free_plan_df.to_excel(writer, sheet_name=RESEARCH_FREE_SHEET, index=False)

    if verbose:
        print(f"[planner] wrote {output_path}")
        print(rotation_df.to_string(index=False))

    # ----- 5. Templates -----
    if generate_templates:
        template_dir = str(cfg.get("TemplatePath", "")).strip()
        template_pattern = str(cfg.get("TemplatePattern", "")).strip()
        for month_str in months:
            template_path = _find_template_for_month(template_dir, template_pattern, month_str)
            if template_path is None:
                print(f"[template] no template found for {month_str}")
                continue
            month_name, yy = _month_label(month_str)
            out_name = f"Stationsplan_{month_name}_{yy}_generated.xlsx"
            out_path = os.path.join(template_dir, out_name)
            _rearrange_template(
                template_path, rotation_df, month_str,
                doctors_df, stations_df, out_path,
            )
            if verbose:
                print(f"[template] generated {out_path}")

    return {
        "rotation": rotation_df,
        "vacation": vacation_df,
        "research_free": research_free_plan_df,
    }


# ----------------------------------------------------------------------
# CLI
# ----------------------------------------------------------------------
if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Standalone rotation planner")
    parser.add_argument("--rules", default="Rules.xlsx")
    parser.add_argument("--out", default="RotationPlan.xlsx")
    parser.add_argument("--start", default=None, help="Start month YYYY-MM")
    parser.add_argument("--months", type=int, default=None)
    parser.add_argument("--min-distinct", type=int, default=None,
                        help="Minimum distinct stations per doctor")
    parser.add_argument("--generate-templates", action="store_true")
    parser.add_argument("--quiet", action="store_true")
    args = parser.parse_args()

    plan_rotation(
        rules_path=args.rules,
        output_path=args.out,
        start_month=args.start,
        n_months=args.months,
        min_distinct=args.min_distinct,
        generate_templates=args.generate_templates if args.generate_templates else None,
        verbose=not args.quiet,
    )